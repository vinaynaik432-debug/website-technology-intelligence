"""One-company-per-row exports. Use the bundled Artifact runtime when available."""
from __future__ import annotations
import csv
import io
import json
import os
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def safe_cell(value):
    if not isinstance(value, str):
        return value
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        value = "'" + value
    # Raw JSON is already split into smaller chunks, so the original remains complete.
    return value if len(value) <= 16000 else value[:15900] + " [See full source JSON columns]"


def csv_bytes(table):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(table["headers"])
    writer.writerows([[safe_cell(v) for v in row] for row in table["rows"]])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


def artifact_runtime():
    bundled = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node"
    node = Path(os.environ.get("ARTIFACT_NODE", str(bundled / "bin/node")))
    modules = Path(os.environ.get("ARTIFACT_MODULES", str(bundled / "node_modules")))
    return (node, modules) if node.exists() and (modules / "@oai/artifact-tool").exists() else None


def write_excel(table, path, preview=False):
    path = Path(path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    runtime = artifact_runtime()
    if runtime:
        node, modules = runtime
        bridge = ROOT / ".tools"
        bridge.mkdir(exist_ok=True)
        link = bridge / "node_modules"
        if not link.exists():
            link.symlink_to(modules, target_is_directory=True)
        shutil.copyfile(ROOT / "scripts/build-workbook.mjs", bridge / "build-workbook.mjs")
        payload = path.with_suffix(".table.json")
        payload.write_text(json.dumps(table, ensure_ascii=False))
        run = subprocess.run([str(node), str(bridge / "build-workbook.mjs"), str(payload), str(path),
                              "preview" if preview else "export"], capture_output=True, text=True, timeout=120)
        if run.returncode:
            raise RuntimeError("Excel export failed. " + run.stderr[-800:])
    else:
        # Portable application fallback on machines without the optional Artifact runtime.
        _portable_excel(table, path)
    return path


def _portable_excel(table, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("Install requirements.txt to enable Excel export on this computer.") from None
    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"
    ws.sheet_view.showGridLines = False
    ws["A2"] = "Website technology intelligence"
    ws["A2"].font = Font(name="Arial", size=14, bold=True, color="183B49")
    ws["A3"] = "Detected = matching signal; Not detected = absent from this scan; Unknown = incomplete or weak evidence."
    ws["A4"] = "DNS records are signals, not proof of active use. Employee counts reflect LinkedIn profiles; offices may be incomplete."
    ws["A5"] = f"Evidence scope: {table['scope']} · Minimum confidence: {table['confidence']} · Collected: {table['created_at']}"
    for r, values in enumerate([table["headers"], *table["rows"]], start=7):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(r, c, safe_cell(value))
            cell.font = Font(name="Arial", size=10, color="203846")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if r == 7:
                cell.fill = PatternFill("solid", fgColor="183B49")
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            elif "checked (UTC)" in table["headers"][c - 1] and isinstance(value, str) and value != "Unknown":
                cell.value = datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
                cell.number_format = "dd mmm yyyy hh:mm"
    last = 7 + len(table["rows"])
    if table["rows"]:
        t = Table(displayName="CompanyIntelligence", ref=f"A7:{get_column_letter(len(table['headers']))}{last}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    ws.freeze_panes = "C8"
    ws.row_dimensions[7].height = 62
    for r in range(8, last + 1):
        ws.row_dimensions[r].height = 42
    for c, title in enumerate(table["headers"], start=1):
        ws.column_dimensions[get_column_letter(c)].width = 26 if c <= table["metadata_count"] else 22
        if title.startswith("Full ") or title == "Description":
            ws.column_dimensions[get_column_letter(c)].width = 45
    if table["rows"] and table["state_count"]:
        a = get_column_letter(table["metadata_count"] + 1)
        b = get_column_letter(table["metadata_count"] + table["state_count"])
        area = f"{a}8:{b}{last}"
        dv = DataValidation(type="list", formula1='"Detected,Not detected,Unknown"')
        ws.add_data_validation(dv)
        dv.add(area)
        for state, color in [("Detected", "E2F1EA"), ("Unknown", "FFF1D6")]:
            ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=['"'+state+'"'], fill=PatternFill("solid",fgColor=color)))
    wb.save(path)
