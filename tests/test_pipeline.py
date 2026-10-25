import json
import tempfile
import unittest
import importlib.util
from pathlib import Path
from unittest.mock import patch
from urllib.error import URLError

from intelligence.apify import Apify, budget_status, reserve, update_cost, run_timeout, collect
from intelligence.export import safe_cell, write_excel
from intelligence.model import domain_name, parse_domains, normalize, catalog, detection_state, select_companies, matrix

ROOT = Path(__file__).resolve().parents[1]


class Domains(unittest.TestCase):
    def test_normalizes_without_collapsing_co_uk(self):
        self.assertEqual(domain_name('https://www.Example.co.uk/about?q=1'), 'example.co.uk')
        self.assertEqual(domain_name('shop.example.co.uk'), 'shop.example.co.uk')
        self.assertTrue(domain_name('bücher.de').startswith('xn--'))

    def test_csv_headers_and_duplicates(self):
        self.assertEqual(parse_domains('company,domain\nN26,n26.com\nN26,https://www.n26.com/en\nHelloFresh,hellofresh.com'), ['n26.com','hellofresh.com'])
        self.assertEqual(parse_domains('n26.com\nhellofresh.com'), ['n26.com','hellofresh.com'])

    def test_private_inputs_and_credentials_rejected(self):
        for value in ['localhost','127.0.0.1','http://127.1','http://0177.0.0.1','http://0x7f.0x1','http://169.254.169.254','https://user:pass@example.com','file:///etc/passwd','company name','example.local','http://example.com:8000']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                domain_name(value)


class Evidence(unittest.TestCase):
    def setUp(self):
        self.report = normalize(ROOT / 'tests/fixtures/basic')
        self.a, self.n, self.u = self.report['companies']

    def test_tri_state_and_scope(self):
        self.assertEqual(detection_state(self.a,name='Example Analytics'), 'Yes')
        self.assertEqual(detection_state(self.n,name='Example Analytics'), 'Unknown')
        self.assertEqual(detection_state(self.u,name='Example Analytics'), 'Unknown')
        self.assertEqual(detection_state(self.a,name='Example Store'), 'No')
        self.assertEqual(detection_state(self.a,name='Example CRM',scope='website'), 'No')
        self.assertEqual(detection_state(self.a,name='Example CRM',scope='dns'), 'Yes')

    def test_confidence_changes_result(self):
        self.assertEqual(detection_state(self.n,name='Example Analytics',confidence=0), 'Yes')

    def test_filters_change_shortlist(self):
        self.assertEqual([c['domain'] for c in select_companies(self.report,tools=['Example Analytics'])],['acme.example.com'])
        self.assertEqual(len(select_companies(self.report,tools=['Example Analytics'],states=['Unknown'])),2)
        self.assertEqual(len(select_companies(self.report,categories=['CRM'],scope='website')),0)
        self.assertEqual(len(select_companies(self.report,country='Germany',min_employees=50)),1)
        self.assertEqual(len(select_companies(self.report,min_employees=0)),2)

    def test_one_row_per_company_and_dynamic_headers(self):
        table=matrix(self.report,raw=True)
        self.assertEqual(len(table['rows']),3)
        self.assertEqual(len(table['headers']),len(set(table['headers'])))
        self.assertTrue(all(len(r)==len(table['headers']) for r in table['rows']))
        self.assertTrue(any(h=='Tool | Analytics | Example Analytics' for h in table['headers']))
        self.assertEqual(self.a['facts']['Founders'],'Unknown')

    def test_raw_json_chunks_preserve_everything(self):
        self.a['raw_company']=[{'long_field':'x'*70000+'😀'*12000+'\uffff','nested':{'zero':0,'false':False}}]
        table=matrix(self.report,raw=True)
        positions=[i for i,h in enumerate(table['headers']) if h.startswith('Full company source JSON')]
        restored=json.loads(''.join(table['rows'][0][i] for i in positions))
        self.assertEqual(restored,self.a['raw_company'])
        self.assertTrue(all(len(table['rows'][0][i])<=12000 for i in positions))
        self.assertTrue(all(len(table['rows'][0][i].encode('utf-16-le'))//2<=12000 for i in positions))

    def test_missing_and_ambiguous_company_data(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)
            for name in ['manifest.json','technology.json','firmographics.json']:
                (p/name).write_text((ROOT/'tests/fixtures/basic'/name).read_text())
            rows=json.loads((p/'firmographics.json').read_text())
            rows.append(dict(rows[0],name='Different candidate'))
            (p/'firmographics.json').write_text(json.dumps(rows))
            r=normalize(p)
            self.assertEqual(r['companies'][0]['facts']['Company match'],'Unknown')
            self.assertEqual(r['companies'][0]['facts']['Employees on LinkedIn'],'Unknown')


class Reliability(unittest.TestCase):
    @unittest.skipUnless(importlib.util.find_spec('openpyxl'),'Install requirements.txt for Excel integration checks')
    def test_formula_like_raw_chunks_are_preserved_as_literal_excel_text(self):
        from openpyxl import load_workbook
        r=normalize(ROOT/'tests/fixtures/basic')
        original=[{'long_field':'='*25000+'😀'*9000+'\uffff'}]
        r['companies'][0]['raw_company']=original
        table=matrix(r,raw=True)
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'source.xlsx';write_excel(table,path)
            sheet=load_workbook(path).active
            indexes=[i for i,h in enumerate(table['headers'],1) if h.startswith('Full company source JSON')]
            restored=json.loads(''.join(sheet.cell(8,i).value or '' for i in indexes))
            self.assertEqual(restored,original)
            self.assertTrue(all(sheet.cell(8,i).data_type=='s' for i in indexes))

    def test_batch_timeout_covers_serial_website_checks(self):
        self.assertEqual(run_timeout('technology',1),300)
        self.assertEqual(run_timeout('technology',7),585)
        self.assertEqual(run_timeout('firmographics',7),300)

    def test_invalid_accepted_cap_stops_before_second_actor(self):
        for accepted in [None, float('nan'), float('inf'), -.5, '0.5', True, 2]:
            with self.subTest(accepted=accepted), tempfile.TemporaryDirectory() as d:
                calls=[]
                def request(api,path,method='GET',body=None,query=None):
                    calls.append((path,method))
                    if method=='GET':return {'data':{'pricingInfos':[{'startedAt':'2020','pricingModel':'PAY_PER_EVENT'}]}}
                    if path.endswith('/abort'):return {}
                    return {'data':{'id':'run','status':'RUNNING','options':{'maxTotalChargeUsd':accepted}}}
                with patch('intelligence.apify.Apify.request',request), self.assertRaisesRegex(RuntimeError,'spending cap'):
                    collect('unit-test-token',['example.com'],Path(d)/'run',Path(d)/'budget.json',1)
                self.assertEqual(len([p for p,m in calls if m=='POST' and p.endswith('/runs')]),1)
                self.assertIn(('/actor-runs/run/abort','POST'),calls)

    def test_monitor_timeout_retains_run_links_and_budget(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            def request(api,path,method='GET',body=None,query=None):
                if method=='GET':return {'data':{'pricingInfos':[{'startedAt':'2020','pricingModel':'PAY_PER_EVENT'}]}}
                return {'data':{'id':'run','status':'RUNNING','options':{'maxTotalChargeUsd':.5}}}
            with patch('intelligence.apify.Apify.request',request), patch('intelligence.apify.time.monotonic',side_effect=[0,1000]), self.assertRaisesRegex(RuntimeError,'Monitoring timed out'):
                collect('unit-test-token',['example.com'],root/'run',root/'budget.json',1)
            manifest=json.loads((root/'run/manifest.json').read_text())
            self.assertEqual(manifest['status'],'Needs attention')
            self.assertTrue(manifest['runs']['technology']['url'].endswith('/run'))
            self.assertEqual(budget_status(root/'budget.json')['reserved_usd'],1)

    def test_budget_survives_restart_and_retains_reservations(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'budget.json'
            reserve(p,'one',3)
            update_cost(p,'one',.02)
            self.assertEqual(budget_status(p)['reserved_usd'],3)
            with self.assertRaises(ValueError): reserve(p,'two',3)
            reserve(p,'three',2)
            self.assertEqual(budget_status(p)['reserved_usd'],5)
            for cap in [0,-1,float('nan'),float('inf')]:
                with self.assertRaises(ValueError): reserve(p,'bad',cap)

    def test_post_never_retried_after_unknown_outcome(self):
        api=Apify('test-token')
        with patch.object(api.opener,'open',side_effect=URLError('network down')) as op:
            with self.assertRaises(RuntimeError): api.request('/acts/test~actor/runs',method='POST',body={})
            self.assertEqual(op.call_count,1)

    def test_spreadsheet_formula_injection_blocked(self):
        for value in ['=HYPERLINK("bad")','+1+1','@SUM(A1)','-cmd','\n=1+1','   =1+1']:
            self.assertTrue(safe_cell(value).startswith("'"))
        self.assertEqual(safe_cell(0),0)
        self.assertEqual(safe_cell(False),False)


if __name__=='__main__':unittest.main()
